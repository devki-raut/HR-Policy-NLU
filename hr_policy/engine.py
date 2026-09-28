"""Dependency-light, extractive retrieval: answers only quote indexed policies."""
import argparse
from collections import Counter
from dataclasses import asdict, dataclass
import hashlib
import json
import math
from pathlib import Path
import re
from hr_policy.processing import clean_pages, process_pdf

CATEGORIES = {
    "leave": {"leave", "vacation", "holiday", "pto", "sick", "parental"},
    "remote_work": {"remote", "hybrid", "home", "office", "wfh"},
    "benefits": {"benefits", "insurance", "medical", "health", "coverage"},
    "expenses": {"expense", "expenses", "reimbursement", "travel", "receipt"},
    "conduct": {"conduct", "harassment", "grievance", "discrimination", "complaint"},
    "onboarding": {"onboarding", "joining", "probation", "orientation"},
}
STOP = set("a an the is are of to for in on i my me do does can how what when where and or with about policy please tell many much get have be it our your".split())

def tokens(text):
    return [t for t in re.findall(r"[a-z0-9]+", text.lower()) if t not in STOP]

@dataclass
class Chunk:
    id: str
    source: str
    page: int
    section: str
    category: str
    text: str

def pages(path):
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return [page.extract_text() or "" for page in PdfReader(path).pages]
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(path)
        lines = []
        for paragraph in doc.paragraphs:
            prefix = "## " if paragraph.style.name.startswith("Heading") else ""
            lines.append(prefix + paragraph.text)
        for table in doc.tables:
            lines.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
        return ["\n".join(lines)]
    return [path.read_text(encoding="utf-8")]

def ingest(folder, output, max_words=180):
    if max_words < 1:
        raise ValueError("max_words must be positive")
    folder, output = Path(folder), Path(output)
    chunks = []
    files = sorted(p for p in folder.rglob("*") if p.suffix.lower() in {".pdf", ".docx", ".txt", ".md"})
    if not files:
        raise ValueError("No PDF, DOCX, TXT or Markdown policies found")
    for path in files:
        extracted = [p.text for p in process_pdf(path).pages] if path.suffix.lower() == ".pdf" else clean_pages(pages(path))
        if not any(p.strip() for p in extracted):
            raise ValueError(f"No text in {path.name}; scanned PDFs require OCR before ingestion")
        section = path.stem
        for page_number, page in enumerate(extracted, 1):
            for block in re.split(r"\n\s*\n|(?=^#{1,6} )", page, flags=re.M):
                lines = block.strip().splitlines()
                if not lines:
                    continue
                if lines[0].startswith("#"):
                    section = lines.pop(0).lstrip("# ")
                words = " ".join(lines).split()
                for start in range(0, len(words), max_words):
                    body = " ".join(words[start:start + max_words])
                    counts = {k: len(set(tokens(section + " " + body)) & v) for k, v in CATEGORIES.items()}
                    category = max(counts, key=counts.get) if any(counts.values()) else "general"
                    source = path.relative_to(folder).as_posix()
                    key = f"{source}:{page_number}:{section}:{start}:{body}"
                    chunks.append(Chunk(hashlib.sha256(key.encode()).hexdigest()[:16], source, page_number, section, category, body))
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_suffix(".tmp")
    temp.write_text(json.dumps({"version": 1, "chunks": [asdict(c) for c in chunks]}, indent=2), encoding="utf-8")
    temp.replace(output)
    return chunks

class PolicyIndex:
    def __init__(self, path):
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
        if payload.get("version") != 1:
            raise ValueError("Unsupported policy index version")
        self.chunks = [Chunk(**c) for c in payload["chunks"]]
        self.terms = [Counter(tokens(c.section + " " + c.text)) for c in self.chunks]
        self.df = Counter(t for terms in self.terms for t in terms)

    def answer(self, question, category=None):
        query = set(tokens(question))
        topics = {k for k, v in CATEGORIES.items() if query & v}
        if len(topics) > 1:
            return "Please ask about one topic at a time: " + ", ".join(sorted(topics)) + "."
        if len(query) < 2:
            return "Please add details—for example, ‘How do I request annual leave?’"
        ranked = []
        average = sum(sum(t.values()) for t in self.terms) / max(len(self.terms), 1)
        for chunk, terms in zip(self.chunks, self.terms):
            if category and chunk.category != category:
                continue
            overlap = query & terms.keys()
            if len(overlap) < 2 or len(overlap) / len(query) < .3:
                continue
            score = sum(math.log(1 + (len(self.chunks) - self.df[t] + .5) / (self.df[t] + .5)) * terms[t] * 2.5 / (terms[t] + 1.5 * (.25 + .75 * sum(terms.values()) / max(average, 1))) for t in overlap)
            ranked.append((score, chunk))
        ranked.sort(key=lambda x: x[0], reverse=True)
        if not ranked:
            return "I couldn’t find enough policy evidence to answer that. Please rephrase or contact HR."
        if len(ranked) > 1 and ranked[1][0] >= ranked[0][0] * .9 and ranked[1][1].section != ranked[0][1].section:
            return "Which policy section do you mean: " + " or ".join(c.section for _, c in ranked[:2]) + "? Please include it in your question."
        chunk = ranked[0][1]
        return f"Relevant policy excerpt (please check that it addresses your question):\n\n{chunk.text}\n\nSource: {chunk.source} — {chunk.section}, page {chunk.page} [chunk {chunk.id}]"

def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    load = sub.add_parser("ingest")
    load.add_argument("folder")
    load.add_argument("--output", default="artifacts/policies.json")
    ask = sub.add_parser("ask")
    ask.add_argument("question")
    ask.add_argument("--index", default="artifacts/policies.json")
    args = parser.parse_args()
    if args.command == "ingest":
        print(f"Indexed {len(ingest(args.folder, args.output))} chunks into {args.output}")
    else:
        print(PolicyIndex(args.index).answer(args.question))

if __name__ == "__main__":
    main()
