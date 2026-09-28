import asyncio
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
from rasa_sdk import Tracker
from rasa_sdk.executor import CollectingDispatcher
from hr_policy.actions import ActionAnswerPolicy
from hr_policy.engine import ingest

class IntegrationTests(unittest.TestCase):
    def test_docx_headings_and_tables(self):
        with tempfile.TemporaryDirectory() as folder:
            doc = Document()
            doc.add_heading('Annual leave', level=1)
            doc.add_paragraph('Annual leave requires manager approval.')
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = 'Leave allowance'
            table.cell(0, 1).text = '20 days'
            doc.save(Path(folder) / 'policy.docx')
            chunks = ingest(folder, Path(folder) / 'index.json')
            self.assertTrue(any('20 days' in c.text for c in chunks))
            self.assertTrue(all(c.section == 'Annual leave' for c in chunks))

    def test_blank_pdf_requires_ocr(self):
        with tempfile.TemporaryDirectory() as folder:
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.write(Path(folder) / 'scan.pdf')
            with self.assertRaisesRegex(ValueError, 'OCR'):
                ingest(folder, Path(folder) / 'index.json')

    def test_rasa_action_returns_citation(self):
        with tempfile.TemporaryDirectory() as folder:
            index = Path(folder) / 'index.json'
            ingest('policies/sample', index)
            dispatcher = CollectingDispatcher()
            tracker = Tracker('test', {}, {'text': 'How many annual leave days?', 'intent': {'name': 'ask_leave'}}, [], False, None, {}, 'action_listen')
            with patch.dict('os.environ', {'POLICY_INDEX': str(index)}):
                events = asyncio.run(ActionAnswerPolicy().run(dispatcher, tracker, {}))
            self.assertEqual(events, [])
            self.assertIn('20 days', dispatcher.messages[0]['text'])
            self.assertIn('Source:', dispatcher.messages[0]['text'])
